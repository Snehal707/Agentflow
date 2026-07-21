"""Rebuild Closing slide layout: CTA above, access codes below, no overlap."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

SRC_CANDIDATES = [
    Path(r"C:\Users\ASUS\Downloads\AgentFlow-Hackathon-Pitch (1).pptx"),
    Path(r"C:\Users\ASUS\Downloads\AgentFlow-Hackathon-Pitch.pptx"),
]
OUT_DOWNLOADS = Path(r"C:\Users\ASUS\Downloads\AgentFlow-Hackathon-Pitch.pptx")
OUT_DOCS = Path(r"C:\Users\ASUS\agent-economy\docs\AgentFlow-Hackathon-Pitch.pptx")
QR = Path(r"C:\Users\ASUS\agent-economy\docs\agentflow-qr.png")
BADGE = Path(r"C:\Users\ASUS\agent-economy\agentflow-frontend\public\media-kit\agentflow-logo-badge.png")

BG = RGBColor(0x07, 0x07, 0x08)
CARD = RGBColor(0x10, 0x10, 0x10)
GOLD = RGBColor(0xF2, 0xCA, 0x50)
MUTED = RGBColor(0xA3, 0xA3, 0xA3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CODES = "P2K5K7RW   ·   4RUYHW5X   ·   S3SQFXYB   ·   9VDFRSRD   ·   F6EBUWQP"


def delete_slide(prs: Presentation, index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    sld_id = list(sld_id_lst)[index]
    rId = sld_id.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sld_id_lst.remove(sld_id)


def slide_title(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    return t
    return ""


def set_run(run, *, size_pt: float, bold: bool = False, color: RGBColor = WHITE, align=None):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"


def add_text(slide, left, top, width, height, text, *, size, bold=False, color=WHITE, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        set_run(r, size_pt=size, bold=bold, color=color)
    return box


def add_card(slide, left, top, width, height, *, line=GOLD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = line
    sh.line.width = Pt(1.5)
    try:
        sh.adjustments[0] = 0.1
    except Exception:
        pass
    return sh


def main() -> None:
    src = next(p for p in SRC_CANDIDATES if p.exists())
    prs = Presentation(str(src))

    # Remove access-codes-only slide if present
    for i in range(len(prs.slides) - 1, -1, -1):
        if slide_title(prs.slides[i]).lower().startswith("access codes"):
            delete_slide(prs, i)

    # Drop the old closing and rebuild it cleanly
    delete_slide(prs, len(prs.slides) - 1)

    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    s = prs.slides.add_slide(blank)

    # Background
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    if BADGE.exists():
        s.shapes.add_picture(str(BADGE), Inches(6.15), Inches(0.22), Inches(1.0), Inches(1.0))

    add_text(
        s, Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.85),
        "Today's AI answers questions.\nAgentFlow completes financial work.",
        size=24, bold=True, color=WHITE,
    )
    add_text(
        s, Inches(0.8), Inches(2.15), Inches(11.7), Inches(0.65),
        "The future isn't opening more apps.\nIt's expressing your intent — on web, Telegram, and soon mobile.",
        size=15, color=MUTED,
    )
    add_text(
        s, Inches(0.8), Inches(2.9), Inches(11.7), Inches(0.4),
        "Intent in. Execution out.",
        size=26, bold=True, color=GOLD,
    )
    add_text(
        s, Inches(0.8), Inches(3.4), Inches(11.7), Inches(0.3),
        "Built by Snehal · Solo founder  ·  X @AgentFlowone",
        size=13, color=GOLD,
    )

    # CTA card — room left below for codes
    add_card(s, Inches(3.55), Inches(3.9), Inches(6.2), Inches(1.55))
    if QR.exists():
        s.shapes.add_picture(str(QR), Inches(3.85), Inches(4.05), Inches(1.25), Inches(1.25))
    add_text(s, Inches(5.35), Inches(4.15), Inches(4.1), Inches(0.35), "Try AgentFlow today", size=17, bold=True, color=GOLD, align=PP_ALIGN.LEFT)
    add_text(s, Inches(5.35), Inches(4.55), Inches(4.1), Inches(0.3), "Scan -> agentflow.one", size=14, color=WHITE, align=PP_ALIGN.LEFT)
    add_text(s, Inches(5.35), Inches(4.9), Inches(4.1), Inches(0.28), "Live on Arc Testnet", size=12, color=MUTED, align=PP_ALIGN.LEFT)

    # Access codes — clearly below CTA
    add_card(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(0.95))
    add_text(
        s, Inches(0.95), Inches(5.85), Inches(11.4), Inches(0.28),
        "Access codes  ·  one per person",
        size=12, bold=True, color=GOLD, align=PP_ALIGN.CENTER,
    )
    add_text(
        s, Inches(0.95), Inches(6.2), Inches(11.4), Inches(0.35),
        CODES,
        size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )

    add_text(
        s, Inches(0.6), Inches(6.95), Inches(10), Inches(0.3),
        "agentflow.one · Arc Testnet",
        size=11, color=MUTED, align=PP_ALIGN.LEFT,
    )

    prs.save(str(OUT_DOWNLOADS))
    prs.save(str(OUT_DOCS))
    print(f"slides={len(prs.slides)}")
    print(f"wrote {OUT_DOWNLOADS}")


if __name__ == "__main__":
    main()
